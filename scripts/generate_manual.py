"""
StockPulse CIPSA - Manual Técnico (10 slides)
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
import os

C = {
    'primary': RGBColor(0x00, 0xD0, 0x84),
    'dark': RGBColor(0x0F, 0x17, 0x2A),
    'surface': RGBColor(0x1E, 0x29, 0x3B),
    'text': RGBColor(0xF0, 0xF4, 0xF8),
    'muted': RGBColor(0x94, 0xA3, 0xB8),
}
W, H = Inches(13.333), Inches(7.5)

def set_bg(slide, color):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = color
    bg.line.fill.background()
    spTree = slide.shapes._spTree
    sp = bg._element
    spTree.remove(sp)
    spTree.insert(2, sp)

def txt(p, text, size_pt, bold=False, color=None, align=None):
    p.text = text
    for run in p.runs:
        run.font.size = Pt(size_pt)
        run.font.bold = bold
        if color:
            run.font.color.rgb = color
        run.font.name = 'Segoe UI'
    if align:
        p.alignment = align

def add_box(prs, left, top, width, height):
    return prs.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))

def slide_title(prs, title, subtitle=""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, C['dark'])
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(0.4), Inches(0.08), Inches(0.8))
    bar.fill.solid()
    bar.fill.fore_color.rgb = C['primary']
    bar.line.fill.background()
    tb = add_box(prs, 0.7, 0.45, 1.5, 0.5)
    txt(tb.text_frame.paragraphs[0], "G360", 20, True, C['primary'])
    tb = add_box(prs, 0.5, 2.5, 12.333, 1.2)
    txt(tb.text_frame.paragraphs[0], title, 48, True, C['text'])
    if subtitle:
        tb = add_box(prs, 0.5, 3.9, 12.333, 0.8)
        txt(tb.text_frame.paragraphs[0], subtitle, 22, False, C['muted'])
    fb = add_box(prs, 0.5, 6.9, 12.333, 0.4)
    txt(fb.text_frame.paragraphs[0], "CIPSA · Intelligence Division", 12, False, C['muted'], PP_ALIGN.RIGHT)
    return slide

def section_slide(prs, title):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, C['dark'])
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(3.0), Inches(2.0), Inches(0.06))
    line.fill.solid()
    line.fill.fore_color.rgb = C['primary']
    line.line.fill.background()
    tb = add_box(prs, 0.5, 2.6, 12.333, 1)
    txt(tb.text_frame.paragraphs[0], title, 40, True, C['text'])
    return slide

def content_slide(prs, title, items):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, C['dark'])
    tb = add_box(prs, 0.5, 0.3, 12.333, 0.6)
    txt(tb.text_frame.paragraphs[0], title, 28, True, C['text'])
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(0.9), Inches(1.5), Inches(0.03))
    line.fill.solid()
    line.fill.fore_color.rgb = C['primary']
    line.line.fill.background()
    cb = add_box(prs, 0.5, 1.2, 12.333, 5.8)
    tf = cb.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
        if item.startswith("##"):
            txt(p, item[2:].strip(), 18, True, C['primary'])
        elif item.startswith("•"):
            txt(p, "▸ " + item[1:].strip(), 14, False, C['text'])
        elif item.startswith("  -"):
            txt(p, "  " + item.strip(), 12, False, C['muted'])
        elif item.startswith("> "):
            txt(p, "「 " + item[2:].strip() + " 」", 12, False, C['primary'])
        elif item == "":
            txt(p, "", 8)
        else:
            txt(p, item, 14, False, C['text'])
    return slide

def image_slide(prs, title, caption=""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, C['dark'])
    tb = add_box(prs, 0.5, 0.3, 12.333, 0.6)
    txt(tb.text_frame.paragraphs[0], title, 28, True, C['text'])
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(0.9), Inches(1.5), Inches(0.03))
    line.fill.solid()
    line.fill.fore_color.rgb = C['primary']
    line.line.fill.background()
    ph = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(1.3), Inches(12.1), Inches(5.5))
    ph.fill.solid()
    ph.fill.fore_color.rgb = C['surface']
    ph.line.color.rgb = C['muted']
    ph.line.width = Pt(1)
    tf = ph.text_frame
    tf.word_wrap = True
    txt(tf.paragraphs[0], "[ CAPTURA DE PANTALLA ]", 22, True, C['muted'], PP_ALIGN.CENTER)
    p2 = tf.add_paragraph()
    txt(p2, caption, 12, False, C['muted'], PP_ALIGN.CENTER)
    return slide

def main():
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    
    # 1: PORTADA
    slide_title(prs, "StockPulse CIPSA", "Manual Técnico — Reportes de Stock en Tiempo Real")
    
    # 2: ARQUITECTURA
    content_slide(prs, "Arquitectura del Sistema", [
        "## Frontend PWA",
        "• Lit 3 Web Components, deploy en GitHub Pages",
        "• Cache 3 capas: memoria → sessionStorage → localStorage",
        "",
        "## Backend API",
        "• FastAPI en Render, fuente: appweb.cipsa.com.pe",
        "• Enriquecimiento automático con catálogo maestro",
        "• TTL cache: 15 minutos",
        "",
        "## Catálogo Maestro",
        "• JSON en GitHub (g360-master-data)",
        "• Campos: SKU, descripción, un_bx, peso, líneas, categorías"
    ])
    
    # 3: FLUJO
    image_slide(prs, "Flujo de Datos", "Captura del diagrama arquitectura")
    
    # 4: CATEGORÍAS
    content_slide(prs, "Categorías de Negocio", [
        "## Mapeo Automático por Línea",
        "• VINIBALL → líneas: 01, MA, 14, AD",
        "• VINIFAN → líneas: 02, 09, 11, 72-79",
        "• REPRESENTADAS → línea: 85",
        "• PUBLICIDAD → líneas: 80, 81",
        "",
        "## Normalización",
        "> Entrada: '0179 - ACCESORIOS' → Código: '79' → VINIFAN",
        "",
        "## Estados",
        "• OK → bx >= 10 | BAJO → 1-9 bx | AGOTADO → bx = 0"
    ])
    
    # 5: FUNCIONALIDADES
    content_slide(prs, "Funcionalidades Principales", [
        "## Dashboard y Estado",
        "• KPIs: total SKUs, cajas, estados por categoría",
        "• Lista expandible por categoría con SKUs detallados",
        "",
        "## Búsqueda Avanzada",
        "• Fuse.js: SKU, nombre, keywords, EAN, línea, categoría",
        "• Búsqueda por voz: normaliza números hablados a SKU",
        "",
        "## Reportes XLSX",
        "• Completo: Resumen + hojas por línea",
        "• Estado: ConStock, BajoStock, SinStock, SinCatálogo"
    ])
    
    # 6: ALERTAS
    content_slide(prs, "Alertas de Stock", [
        "## Tipos de Alerta",
        "• Crítico (rojo): SKU agotado (bx = 0)",
        "• Advertencia (amarillo): Stock bajo (1-9 cajas)",
        "",
        "## Criterios",
        "• Solo SKUs con estado_linea definido (catálogo)",
        "• Ordenados por prioridad y cantidad de cajas",
        "",
        "## Visualización",
        "• Badge con contador en navegación",
        "• Panel dedicado con filtros por tipo"
    ])
    
    # 7: SEGURIDAD
    content_slide(prs, "Seguridad y Acceso", [
        "## Modelo de Autenticación",
        "• S1_API_KEY: Administrativa (upload, catálogo)",
        "• S1_READ_API_KEY: Lectura para frontend",
        "",
        "## Endpoints Protegidos",
        "• GET /stock, /health → requiere READ_KEY",
        "• POST /catalog/upload → requiere ADMIN_KEY",
        "",
        "## Controles",
        "• CORS: github.io + localhost",
        "• Rate limit: 60 req/min por IP",
        "• Ventana operativa: Lun-Sáb 07:00-22:59"
    ])
    
    # 8: DESPLIEGUE
    content_slide(prs, "Despliegue y Operación", [
        "## Infraestructura",
        "• Frontend: GitHub Pages (auto-deploy desde main)",
        "• Backend: Render (gratuito, keep-alive cada 5 min)",
        "",
        "## Desarrollo Local",
        "• Frontend: npm run dev (localhost:3000)",
        "• Backend: uvicorn app.main:app --reload",
        "",
        "## Variables Clave",
        "• S1_SOURCE1_URL, S1_SOURCE2_URL (appweb)",
        "• S1_API_KEY, S1_READ_API_KEY (separar permisos)"
    ])
    
    # 9: TROUBLESHOOTING
    content_slide(prs, "Troubleshooting Común", [
        "## Cache Stale",
        "• Síntoma: datos no se actualizan",
        "• Solución: Forzar refresh o limpiar localStorage",
        "",
        "## Categoría OTROS",
        "• Causa: código de línea no mapeado",
        "• Solución: Agregar a CATEGORIAS en constants.py",
        "",
        "## Stock 0 con existencias",
        "• Causa: stock solo en almacén mktd (118, S*)",
        "• Solución: Verificar tipo de almacén en datos"
    ])
    
    # 10: CIERRE
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, C['dark'])
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(3.0), Inches(2.0), Inches(0.06))
    line.fill.solid()
    line.fill.fore_color.rgb = C['primary']
    line.line.fill.background()
    tb = add_box(prs, 0.5, 2.0, 12.333, 1.0)
    txt(tb.text_frame.paragraphs[0], "StockPulse CIPSA", 44, True, C['text'], PP_ALIGN.CENTER)
    sb = add_box(prs, 0.5, 3.3, 12.333, 0.6)
    txt(sb.text_frame.paragraphs[0], "Inteligencia de Stock en Tiempo Real", 20, False, C['muted'], PP_ALIGN.CENTER)
    cb = add_box(prs, 0.5, 5.0, 12.333, 0.5)
    txt(cb.text_frame.paragraphs[0], "¿Consultas? Contactar al equipo G360", 14, False, C['primary'], PP_ALIGN.CENTER)
    fb = add_box(prs, 0.5, 6.8, 12.333, 0.4)
    txt(fb.text_frame.paragraphs[0], "g360-stock-reporter-lit · GitHub", 10, False, C['muted'], PP_ALIGN.RIGHT)
    
    output_path = os.path.join(os.path.dirname(__file__), "StockPulse_Manual_Tecnico.pptx")
    prs.save(output_path)
    print(f"[OK] Manual: {len(prs.slides)} slides")

if __name__ == "__main__":
    main()
